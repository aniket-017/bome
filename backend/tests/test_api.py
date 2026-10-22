"""API tests against the real database (skipped when PostgreSQL is not reachable)."""
import io

import pytest
from openpyxl import Workbook, load_workbook

try:
    from fastapi.testclient import TestClient
    from sqlalchemy import text

    from app.db import engine
    from app.main import app

    with engine.connect() as conn:
        HAS_PARTS = conn.execute(text("SELECT count(*) FROM parts")).scalar() > 0
except Exception:  # pragma: no cover - environment without DB
    HAS_PARTS = False

pytestmark = pytest.mark.skipif(not HAS_PARTS, reason="database with parts master required")


REDUCER = "REDUCER CONCENTRIC SS ASTM A403 WP304 DN 150 X DN 100 SCH 40S BW ASME B16.9"


@pytest.fixture(scope="module")
def client():
    return TestClient(app)


def item_no_of(description: str) -> str:
    with engine.connect() as conn:
        return conn.execute(text("SELECT item_no FROM parts WHERE description = :d"), {"d": description}).scalar_one()


def _bom_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["PROJECT X"])
    ws.append([])
    ws.append(["Sr No", "Description", "Qty", "UOM"])
    ws.append([1, "Conc Red. A403 WP304 6\" x 4\" BW Sch 40", 2, "NOS"])
    ws.append([2, "Hex bolt IS 1367 Gr 8.8 M24 x 100 HDG full thread", 16, "NOS"])
    ws.append([3, "Lifting lug as per drawing", 4, "NOS"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_search(client):
    r = client.get("/api/parts/search", params={"q": "stud m20 x 310", "limit": 5})
    assert r.status_code == 200
    body = r.json()
    assert body["items"] and all("M20" in i["description"] for i in body["items"])


def test_search_by_noun_and_browse_by_modifier(client):
    body = client.get("/api/parts/search", params={"q": "6205", "noun": "BEARING", "limit": 5}).json()
    assert body["items"] and all(i["noun"] == "BEARING" for i in body["items"])

    nouns = client.get("/api/parts/nouns", params={"group": "FLANGES"}).json()
    assert {"FLANGE", "SPECTACLE BLIND"} <= {n["noun"] for n in nouns}

    flange = client.get("/api/parts/nouns/FLANGE").json()
    assert [m["modifier"] for m in flange["modifiers"]] == [
        m for m in flange["modifier_order"] if m in {x["modifier"] for x in flange["modifiers"]}]
    assert "WN" in {v["value"] for v in flange["modifiers"][0]["values"]}

    r = client.get("/api/parts/browse", params=[("noun", "FLANGE"), ("m", "TYPE=WN"), ("m", "RATING=CL 300"),
                                                 ("m", "SIZE=DN 100"), ("limit", 5)]).json()
    assert r["total"] > 0 and len(r["items"]) <= 5
    assert all(i["modifiers"]["TYPE"] == "WN" and i["modifiers"]["RATING"] == "CL 300" for i in r["items"])
    assert client.get("/api/parts/browse", params={"noun": "FLANGE", "m": "TYPE"}).status_code == 422

    part = client.get(f"/api/parts/{r['items'][0]['id']}").json()
    assert part["noun"] == "FLANGE" and part["modifiers"]["SIZE"] == "DN 100"


def test_excel_upload_flow(client):
    data = _bom_xlsx()
    files = {"file": ("bom.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post("/api/boms/upload/preview", files=files).json()
    assert preview["header_row"] == 2
    assert preview["suggested_mapping"] == {"line_no": 0, "description": 1, "qty": 2, "uom": 3}

    r = client.post(
        "/api/boms/upload",
        files={"file": ("bom.xlsx", data)},
        data={"name": "pytest upload", "header_row": "2", "mapping": '{"line_no":0,"description":1,"qty":2,"uom":3}'},
    )
    assert r.status_code == 201, r.text
    bom = r.json()
    try:
        statuses = [l["match_status"] for l in bom["lines"]]
        assert statuses[0] == "auto" and statuses[1] == "auto"
        assert bom["lines"][0]["matched_part"]["item_no"] == item_no_of(REDUCER)
        assert statuses[2] in ("unmatched", "review")
        assert bom["status"] == "needs_review"

        # manual override of the unmatched line, then the BOM is fully matched
        line3 = bom["lines"][2]
        r = client.patch(f"/api/boms/{bom['id']}/lines/{line3['id']}", json={"matched_part_id": bom["lines"][0]["matched_part_id"]})
        assert r.status_code == 200 and r.json()["match_status"] == "manual"
        assert client.get(f"/api/boms/{bom['id']}").json()["status"] == "matched"

        export = client.get(f"/api/boms/{bom['id']}/export")
        assert export.status_code == 200
        ws = load_workbook(io.BytesIO(export.content)).active
        assert ws.cell(row=2, column=8).value == item_no_of(REDUCER)
    finally:
        client.delete(f"/api/boms/{bom['id']}")


def test_manual_bom_and_line_edit(client):
    r = client.post("/api/boms", json={"name": "pytest manual", "lines": [
        {"raw_description": "Stud full thread A193 B7 M20 x 300 plain", "qty": 8},
    ]})
    assert r.status_code == 201, r.text
    bom = r.json()
    try:
        line = bom["lines"][0]
        assert line["match_status"] in ("auto", "exact_desc")
        r = client.patch(f"/api/boms/{bom['id']}/lines/{line['id']}",
                         json={"raw_description": "PLATE CS SA-516 GR 70 THK- 12 mm 2500 x 6300 mm"})
        assert r.json()["match_status"] == "exact_desc"
        r = client.post(f"/api/boms/{bom['id']}/lines", json={"raw_description": "elbow"})
        assert r.status_code == 201
        assert r.json()["line_no"] == 2
    finally:
        client.delete(f"/api/boms/{bom['id']}")
