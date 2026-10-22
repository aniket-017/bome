import io
import json
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.bom_service import (
    FIELDS,
    PREVIEW_ROWS,
    apply_result,
    bom_to_dict,
    create_bom,
    json_cell,
    load_bom,
    match_bom_lines,
    part_exists,
    refresh_status,
    rows_to_lines,
    suggest_layout,
)
from app.config import get_settings
from app.db import get_db
from app.matcher import match_lines
from app.models import Bom, BomLine
from app.schemas import AcceptRequest, BomCreate, BomLineIn, BomLineOut, BomLineUpdate, BomOut, BomSummary, MatchRequest
from app.tabular import list_sheets, read_rows
from app.vision import DOCUMENT_MIME_TYPES, VisionError, VisionNotConfigured, extract_bom_lines

router = APIRouter(prefix="/api/boms", tags=["boms"])

MAX_UPLOAD_BYTES = 30 * 1024 * 1024
TABULAR_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".xlsb", ".ods", ".csv")
IMAGE_SUFFIXES = tuple(DOCUMENT_MIME_TYPES)


async def _read_upload(file: UploadFile, suffixes: tuple[str, ...]) -> bytes:
    name = (file.filename or "").lower()
    if not name.endswith(suffixes):
        raise HTTPException(415, f"Unsupported file type. Allowed: {', '.join(suffixes)}")
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "File too large (max 30 MB)")
    if not data:
        raise HTTPException(422, "File is empty")
    return data


def _read_table(data: bytes, filename: str, sheet: str | None) -> list[list]:
    try:
        return read_rows(data, filename, sheet or None)
    except Exception as exc:
        raise HTTPException(422, f"Could not read the file: {exc}") from exc


@router.get("", response_model=list[BomSummary])
def list_boms(db: Session = Depends(get_db)):
    boms = db.scalars(select(Bom).order_by(Bom.created_at.desc()).limit(500)).all()
    counts: dict[int, dict[str, int]] = {}
    if boms:
        rows = db.execute(
            select(BomLine.bom_id, BomLine.match_status, func.count())
            .where(BomLine.bom_id.in_([b.id for b in boms]))
            .group_by(BomLine.bom_id, BomLine.match_status)
        )
        for bom_id, status, n in rows:
            counts.setdefault(bom_id, {})[status] = n
    return [
        {
            "id": b.id, "name": b.name, "project": b.project, "source_type": b.source_type,
            "source_file": b.source_file, "status": b.status, "created_at": b.created_at,
            "line_count": sum(counts.get(b.id, {}).values()), "status_counts": counts.get(b.id, {}),
        }
        for b in boms
    ]


@router.post("", response_model=BomOut, status_code=201)
def create_manual_bom(body: BomCreate, db: Session = Depends(get_db)):
    for line in body.lines:
        if line.matched_part_id:
            part_exists(db, line.matched_part_id)
    bom = create_bom(
        db, name=body.name, project=body.project, source_type="manual",
        lines=[l.model_dump() for l in body.lines], auto_match=body.auto_match,
    )
    return bom_to_dict(bom)


@router.post("/upload/preview")
async def upload_preview(
    file: UploadFile = File(...),
    sheet: str | None = Form(None),
    header_row: int | None = Form(None),
):
    data = await _read_upload(file, TABULAR_SUFFIXES)
    filename = file.filename or "upload.xlsx"
    sheets = list_sheets(data, filename)
    sheet_name = sheet if sheet in sheets else sheets[0]
    rows = _read_table(data, filename, None if filename.lower().endswith(".csv") else sheet_name)
    header_idx, mapping = suggest_layout(rows, header_row)
    header = rows[header_idx] if header_idx < len(rows) else []
    body = rows[header_idx + 1:]
    width = max((len(r) for r in rows[header_idx: header_idx + PREVIEW_ROWS + 1]), default=0)
    return {
        "sheets": sheets,
        "sheet": sheet_name,
        "header_row": header_idx,
        "headers": [str(json_cell(h)) if h not in (None, "") else f"Column {i + 1}" for i, h in
                    enumerate(list(header) + [None] * (width - len(header)))],
        "rows": [[json_cell(v) if v != "" else None for v in r] for r in body[:PREVIEW_ROWS]],
        "total_rows": sum(1 for r in body if any(v not in (None, "") for v in r)),
        "suggested_mapping": mapping,
    }


@router.post("/upload", response_model=BomOut, status_code=201)
async def upload_bom(
    file: UploadFile = File(...),
    name: str = Form(...),
    project: str | None = Form(None),
    sheet: str | None = Form(None),
    header_row: int = Form(0),
    mapping: str = Form(...),
    db: Session = Depends(get_db),
):
    data = await _read_upload(file, TABULAR_SUFFIXES)
    try:
        col_map = {k: int(v) for k, v in json.loads(mapping).items() if k in FIELDS and v is not None and v != ""}
    except (ValueError, AttributeError) as exc:
        raise HTTPException(422, "mapping must be a JSON object of field -> column index") from exc
    filename = file.filename or "upload.xlsx"
    rows = _read_table(data, filename, sheet)
    if header_row >= len(rows):
        raise HTTPException(422, "header_row is beyond the end of the sheet")
    lines = rows_to_lines(rows, header_row, col_map)
    if not lines:
        raise HTTPException(422, "No BOM lines found below the header row")
    bom = create_bom(db, name=name, project=project, source_type="excel", lines=lines, source_file=filename)
    return bom_to_dict(bom)


@router.post("/upload-image", response_model=BomOut, status_code=201)
async def upload_bom_image(
    file: UploadFile = File(...),
    name: str = Form(...),
    project: str | None = Form(None),
    db: Session = Depends(get_db),
):
    data = await _read_upload(file, IMAGE_SUFFIXES)
    suffix = Path(file.filename or "").suffix.lower()
    try:
        lines = await run_in_threadpool(extract_bom_lines, data, DOCUMENT_MIME_TYPES[suffix])
    except VisionNotConfigured as exc:
        raise HTTPException(503, str(exc)) from exc
    except VisionError as exc:
        raise HTTPException(502, str(exc)) from exc
    if not lines:
        raise HTTPException(422, "No BOM lines could be read from the image")

    path = get_settings().upload_path / f"{uuid.uuid4().hex}{suffix}"
    path.write_bytes(data)
    try:
        bom = create_bom(db, name=name, project=project, source_type="image", lines=lines,
                         source_file=file.filename, image_path=str(path))
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return bom_to_dict(bom)


@router.get("/{bom_id}", response_model=BomOut)
def get_bom(bom_id: int, db: Session = Depends(get_db)):
    return bom_to_dict(load_bom(db, bom_id))


@router.delete("/{bom_id}", status_code=204)
def delete_bom(bom_id: int, db: Session = Depends(get_db)):
    bom = load_bom(db, bom_id)
    if bom.image_path:
        Path(bom.image_path).unlink(missing_ok=True)
    db.delete(bom)
    db.commit()


@router.post("/{bom_id}/match", response_model=BomOut)
def rematch_bom(bom_id: int, body: MatchRequest | None = None, db: Session = Depends(get_db)):
    force = bool(body and body.force)
    bom = load_bom(db, bom_id)
    targets = [l for l in bom.lines if force or l.match_status != "manual"]
    if targets:
        match_bom_lines(targets)
    refresh_status(bom)
    db.commit()
    return bom_to_dict(load_bom(db, bom_id))


@router.post("/{bom_id}/accept", response_model=BomOut)
def accept_suggestions(bom_id: int, body: AcceptRequest, db: Session = Depends(get_db)):
    """Accept the top candidate for review lines (or for the given line ids)."""
    bom = load_bom(db, bom_id)
    wanted = set(body.line_ids) if body.line_ids else None
    for line in bom.lines:
        if wanted is not None and line.id not in wanted:
            continue
        if wanted is None and line.match_status != "review":
            continue
        if not line.candidates or line.candidates[0]["score"] < body.min_score:
            continue
        top = line.candidates[0]
        line.matched_part_id = top["part_id"]
        line.match_score = top["score"]
        line.match_status = "manual"
    refresh_status(bom)
    db.commit()
    return bom_to_dict(load_bom(db, bom_id))


def _get_line(db: Session, bom_id: int, line_id: int) -> BomLine:
    line = db.get(BomLine, line_id)
    if not line or line.bom_id != bom_id:
        raise HTTPException(404, "BOM line not found")
    return line


@router.post("/{bom_id}/lines", response_model=BomLineOut, status_code=201)
def add_line(bom_id: int, body: BomLineIn, db: Session = Depends(get_db)):
    bom = load_bom(db, bom_id)
    line = BomLine(
        line_no=body.line_no or (max((l.line_no for l in bom.lines), default=0) + 1),
        raw_item_no=body.raw_item_no, raw_description=body.raw_description, qty=body.qty, uom=body.uom,
        candidates=[], extra={},
    )
    if body.matched_part_id:
        part_exists(db, body.matched_part_id)
        line.matched_part_id = body.matched_part_id
        line.match_status = "manual"
    else:
        apply_result(line, match_lines(db, [(body.raw_item_no, body.raw_description)])[0])
    bom.lines.append(line)
    refresh_status(bom)
    db.commit()
    db.refresh(line)
    return line


@router.patch("/{bom_id}/lines/{line_id}", response_model=BomLineOut)
def update_line(bom_id: int, line_id: int, body: BomLineUpdate, db: Session = Depends(get_db)):
    line = _get_line(db, bom_id, line_id)
    changes = body.model_dump(exclude_unset=True)
    text_changed = False
    for field in ("raw_item_no", "raw_description", "qty", "uom"):
        if field in changes:
            text_changed |= field in ("raw_item_no", "raw_description") and getattr(line, field) != changes[field]
            setattr(line, field, changes[field])

    if body.matched_part_id:
        part_exists(db, body.matched_part_id)
        line.matched_part_id = body.matched_part_id
        line.match_status = "manual"
        cand = next((c for c in line.candidates if c["part_id"] == body.matched_part_id), None)
        line.match_score = cand["score"] if cand else None
    elif body.clear_match:
        line.matched_part_id = None
        line.match_score = None
        line.match_status = "unmatched"
    elif body.rematch or text_changed:
        apply_result(line, match_lines(db, [(line.raw_item_no, line.raw_description)])[0])

    db.flush()
    refresh_status(load_bom(db, bom_id))
    db.commit()
    db.refresh(line)
    return line


@router.delete("/{bom_id}/lines/{line_id}", status_code=204)
def delete_line(bom_id: int, line_id: int, db: Session = Depends(get_db)):
    line = _get_line(db, bom_id, line_id)
    db.delete(line)
    db.flush()
    refresh_status(load_bom(db, bom_id))
    db.commit()


STATUS_FILLS = {
    "exact_item_no": "C6EFCE", "exact_desc": "C6EFCE", "auto": "DDEBF7",
    "review": "FFEB9C", "unmatched": "FFC7CE", "manual": "E4DFEC",
}


@router.get("/{bom_id}/export")
def export_bom(bom_id: int, db: Session = Depends(get_db)):
    bom = load_bom(db, bom_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Matched BOM"
    extra_cols: list[str] = []
    for line in bom.lines:
        for key in (line.extra or {}).get("row", {}):
            if key not in extra_cols:
                extra_cols.append(key)
    headers = ["Line", "BOM Item No", "BOM Description", "Qty", "UoM", "Match Status", "Score",
               "Matched Item No", "Matched Description", "Item Group", "Top Suggestion Item No",
               "Top Suggestion Description"] + [f"Source: {c}" for c in extra_cols]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")

    for line in bom.lines:
        part = line.matched_part
        top = line.candidates[0] if line.candidates and not part else None
        row_data = (line.extra or {}).get("row", {})
        ws.append([
            line.line_no, line.raw_item_no, line.raw_description,
            float(line.qty) if line.qty is not None else None, line.uom,
            line.match_status, line.match_score,
            part.item_no if part else None, part.description if part else None, part.item_group if part else None,
            top["item_no"] if top else None, top["description"] if top else None,
            *[row_data.get(c) for c in extra_cols],
        ])
        ws.cell(row=ws.max_row, column=6).fill = PatternFill("solid", fgColor=STATUS_FILLS.get(line.match_status, "FFFFFF"))

    widths = {1: 6, 2: 14, 3: 60, 4: 8, 5: 7, 6: 14, 7: 7, 8: 14, 9: 60, 10: 14, 11: 14, 12: 60}
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 18)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in bom.name).strip() or f"bom_{bom.id}"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe}_matched.xlsx"'},
    )


@router.get("/{bom_id}/image")
def get_bom_image(bom_id: int, db: Session = Depends(get_db)):
    bom = load_bom(db, bom_id)
    if not bom.image_path:
        raise HTTPException(404, "This BOM has no source image")
    path = Path(bom.image_path)
    upload_root = get_settings().upload_path.resolve()
    if not path.resolve().is_relative_to(upload_root) or not path.exists():
        raise HTTPException(404, "Image file missing")
    return FileResponse(path)
